# -*- coding: utf-8 -*-
"""Estimate token count for one criterion evaluation call."""
import sys, os
sys.path.insert(0, r"e:\Phase 1_ Quantum Foundation Project Setup Instructions\backend")
os.chdir(r"e:\Phase 1_ Quantum Foundation Project Setup Instructions\backend")

import docx, re

def extract_full(path):
    doc = docx.Document(path)
    lines = []
    for p in doc.paragraphs:
        if p.text.strip(): lines.append(p.text.strip())
    seen = set(lines)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if t and t not in seen:
                    lines.append(t)
                    seen.add(t)
    return '\n'.join(lines)

brief   = extract_full(r"D:\OneDrive\Desktop\Marj Alhamam secondary..ID 96294 -L3-Y2\Official Assessment  BUS - L3 - Y2 U14 - Part 1\Official Assessment  BUS - L3 - Y2 U14 - Part 1\L3 BUS U14 AAB Learning Aim AB V02 Jun-24 (Arabic).docx")
student = extract_full(r"D:\OneDrive\Desktop\Marj Alhamam secondary..ID 96294 -L3-Y2\ALL\AMIR ALHERBAWI\خدمه عملاء تعديل2.docx")

from app.archive.forensic_engine import extract_relevant_excerpt, sample_document

# Test excerpt size
excerpt = extract_relevant_excerpt(student, "خدمة العملاء PESTLE", max_chars=14000)
print(f"Brief: {len(brief):,} chars")
print(f"Student full: {len(student):,} chars")
print(f"Student excerpt (14k): {len(excerpt):,} chars")

# Rough token estimate (Arabic: 1 char ≈ 1 token, ASCII: 4 chars ≈ 1 token)
def count_tokens_approx(text):
    arabic = sum(1 for c in text if '\u0600' <= c <= '\u06FF')
    return arabic + (len(text) - arabic) // 4

print(f"\nToken estimates:")
print(f"  Brief: ~{count_tokens_approx(brief):,} tokens")
print(f"  Student excerpt: ~{count_tokens_approx(excerpt):,} tokens")

# Estimate the template overhead (system instructions text in the prompt)
template_overhead = """أنت مقيّم أكاديمي متخصص في مؤهلات BTEC
PASS/MERIT/DISTINCTION definitions + steps + JSON format"""
print(f"\n  Template overhead: estimate ~5,000-8,000 tokens")
print(f"\n  TOTAL estimate: ~{count_tokens_approx(brief) + count_tokens_approx(excerpt) + 6000:,} tokens")
print(f"\nOpenAI TPM limit: 30,000 tokens")
print(f"Within limit? {(count_tokens_approx(brief) + count_tokens_approx(excerpt) + 6000) < 30000}")
