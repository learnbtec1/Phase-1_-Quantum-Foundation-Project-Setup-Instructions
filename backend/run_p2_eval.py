# -*- coding: utf-8 -*-
import docx, sys, json, requests
sys.stdout.reconfigure(encoding='utf-8')

# ---- Extract texts ----
def extract_text(path):
    doc = docx.Document(path)
    lines = []
    for p in doc.paragraphs:
        if p.text.strip():
            lines.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if t and t not in lines:
                    lines.append(t)
    return '\n'.join(lines)

brief = extract_text(r"D:\OneDrive\Desktop\Marj Alhamam secondary..ID 96294 -L3-Y2\Official Assessment  BUS - L3 - Y2 U14 - Part 1\Official Assessment  BUS - L3 - Y2 U14 - Part 1\L3 BUS U14 AAB Learning Aim AB V02 Jun-24 (Arabic).docx")
amir  = extract_text(r"D:\OneDrive\Desktop\Marj Alhamam secondary..ID 96294 -L3-Y2\ALL\AMIR ALHERBAWI\خدمه عملاء تعديل2.docx")

# Build brief with ONLY A.P2 criterion — no mention of 'شركتين' to avoid false global quantitative constraint.
# A.P2 explicitly says 'الشركة المختارة' (the chosen company — singular).
# We include a dummy A.P1 so the backend's minimum-2-criteria check passes.
p2_brief = """موجز واجب BTEC
الوحدة 14: دراسة خدمة العملاء

المعايير:
14/A.P1
وصف الطرق المختلفة لتقديم خدمة العملاء في الشركة المختارة.

14/A.P2
دراسه( فحص ) الطرق التي يمكن أن تستخدمها خدمة العملاء في الشركة المختارة لتلبية توقعات العملاء وتحقيق رضاهم والالتزام بالتشريعات واللوائح الحالية ذات الصلة.
"""

print(f"Brief length: {len(p2_brief)} | Amir text length: {len(amir)}", flush=True)
print("Sending to backend (P2 only)...", flush=True)

# ---- POST to backend ----
payload = {
    "assignment_text": p2_brief,
    "student_text": amir
}

try:
    resp = requests.post(
        "http://127.0.0.1:8000/api/v1/assessment/forensic-grade-v3",
        json=payload,
        timeout=300
    )
    print(f"HTTP {resp.status_code}", flush=True)
except Exception as e:
    print(f"REQUEST FAILED: {e}", flush=True)
    raise
if resp.status_code == 200:
    data = resp.json()
    # Print final grade
    print(f"\n{'='*60}")
    print(f"FINAL GRADE: {data.get('final_grade', 'N/A')}")
    print(f"{'='*60}")
    # Print A.P2 specifically
    criteria = data.get('criteria', {})
    for code, detail in criteria.items():
        if 'P2' in code or 'p2' in code.lower():
            print(f"\n--- {code} ---")
            print(f"  Achieved  : {detail.get('achieved')}")
            print(f"  Band      : {detail.get('band')}")
            print(f"  Feedback  : {detail.get('feedback','')}")
            print(f"  Evidence  : {detail.get('evidence_quote','')[:300]}")
    print("\n\n--- ALL CRITERIA ---")
    for code, detail in criteria.items():
        achieved = "✅" if detail.get('achieved') else "❌"
        print(f"  {achieved} {code}: {detail.get('feedback','')[:120]}")
    print(f"\nSummary: {data.get('summary','')[:500]}")
else:
    print("ERROR:", resp.text[:2000])
