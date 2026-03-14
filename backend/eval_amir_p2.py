# -*- coding: utf-8 -*-
import docx, json, asyncio, sys, os
sys.path.insert(0, r"e:\Phase 1_ Quantum Foundation Project Setup Instructions\backend")
os.chdir(r"e:\Phase 1_ Quantum Foundation Project Setup Instructions\backend")

# Read assignment brief — extract ALL text including tables
def extract_full_text(doc_path):
    doc = docx.Document(doc_path)
    lines = []
    for p in doc.paragraphs:
        if p.text.strip():
            lines.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text and cell_text not in lines:
                    lines.append(cell_text)
    return '\n'.join(lines)

brief_text = extract_full_text(r"D:\OneDrive\Desktop\Marj Alhamam secondary..ID 96294 -L3-Y2\Official Assessment  BUS - L3 - Y2 U14 - Part 1\Official Assessment  BUS - L3 - Y2 U14 - Part 1\L3 BUS U14 AAB Learning Aim AB V02 Jun-24 (Arabic).docx")

# Read Amir's answer
amir_doc = docx.Document(r"D:\OneDrive\Desktop\Marj Alhamam secondary..ID 96294 -L3-Y2\ALL\AMIR ALHERBAWI\خدمه عملاء تعديل2.docx")
amir_text = '\n'.join([p.text for p in amir_doc.paragraphs if p.text.strip()])

print("=== BRIEF (first 3000 chars) ===")
print(brief_text[:3000])
print("\n\n=== P2 in BRIEF ===")
# Find P2 description in brief
import re
p2_match = re.search(r'A\.P2.{0,1000}', brief_text, re.DOTALL)
if p2_match:
    print(p2_match.group(0)[:800])
else:
    print("-- P2 pattern not found, searching 'P2' --")
    idx = brief_text.find('P2')
    if idx >= 0:
        print(brief_text[idx:idx+600])

print("\n\n=== AMIR P2 SECTION ===")
idx = amir_text.find('P2')
if idx >= 0:
    print(amir_text[idx:idx+5000])
else:
    print("P2 label not found, showing chars 2000-8000:")
    print(amir_text[2000:8000])
