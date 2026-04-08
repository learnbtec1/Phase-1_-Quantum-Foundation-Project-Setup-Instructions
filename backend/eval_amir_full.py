# -*- coding: utf-8 -*-
"""
Amir full evaluation script.
"""
import sys, os
import io
# Force UTF-8 output on Windows
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')
import docx, json, asyncio

sys.path.insert(0, r"e:\Phase 1_ Quantum Foundation Project Setup Instructions\backend")
os.chdir(r"e:\Phase 1_ Quantum Foundation Project Setup Instructions\backend")

BRIEF_PATH  = r"D:\OneDrive\Desktop\Marj Alhamam secondary..ID 96294 -L3-Y2\Official Assessment  BUS - L3 - Y2 U14 - Part 1\Official Assessment  BUS - L3 - Y2 U14 - Part 1\L3 BUS U14 AAB Learning Aim AB V02 Jun-24 (Arabic).docx"
STUDENT_PATH = r"D:\OneDrive\Desktop\Marj Alhamam secondary..ID 96294 -L3-Y2\ALL\AMIR ALHERBAWI\خدمه عملاء تعديل2.docx"

# ─── Extract full text from docx including tables ──────────────────────
def extract_full(path):
    doc = docx.Document(path)
    lines = []
    for p in doc.paragraphs:
        if p.text.strip():
            lines.append(p.text.strip())
    seen = set(lines)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if t and t not in seen:
                    lines.append(t)
                    seen.add(t)
    return '\n'.join(lines)

print("[INFO] Reading files...", flush=True)
brief_text   = extract_full(BRIEF_PATH)
student_text = extract_full(STUDENT_PATH)

print(f"[INFO] Brief: {len(brief_text):,} chars", flush=True)
print(f"[INFO] Amir answer: {len(student_text):,} chars", flush=True)

# ─── Run the full forensic grader ──────────────────────────────────────
from app.archive.forensic_engine import forensic_grade

async def run():
    print("\n[INFO] Running full evaluation (may take 2-4 minutes)...\n", flush=True)
    result = await forensic_grade(brief_text, student_text)
    return result

result = asyncio.run(run())

# ─── Print results ──────────────────────────────────────────────────────
print("\n" + "="*60)
print(f"FINAL GRADE: {result.get('final_grade', 'unknown')}")
print("="*60)
print(f"\nSUMMARY:\n{result.get('summary','')}\n")

criteria = result.get('criteria') or result.get('criteria_results', {})
print("-"*60)
print("CRITERIA DETAILS:\n")

for code, detail in sorted(criteria.items()):
    achieved = detail.get('achieved', False)
    status = "ACHIEVED" if achieved else "NOT ACHIEVED"
    band = detail.get('band', '')
    feedback = detail.get('feedback', '')
    evidence = detail.get('evidence_quote', '')
    conf = detail.get('confidence', '')
    
    print(f"[{status}] [{band}] {code}")
    print(f"   Feedback: {feedback}")
    if evidence:
        preview = evidence[:200].replace('\n', ' ')
        print(f"   Evidence: {preview}{'...' if len(evidence)>200 else ''}")
    if conf:
        print(f"   Confidence: {conf}")
    print()

# Save full result to JSON for reference
with open("amir_eval_result.json", "w", encoding="utf-8") as f:
    json.dump(result, f, ensure_ascii=False, indent=2)
print(f"\nSaved full result to: amir_eval_result.json")
