# -*- coding: utf-8 -*-
"""
file_extractor.py
-----------------
استخراج النص من ملفات .docx / .pptx / .pdf / .txt
يُرجع النص المدمج مع فواصل واضحة بين الملفات.
"""

from __future__ import annotations
import io
import logging
from typing import List, Tuple

logger = logging.getLogger(__name__)


# ─────────────────────────────────────────────
# Individual extractors
# ─────────────────────────────────────────────

def extract_docx(data: bytes) -> str:
    """Extract text from a .docx binary blob."""
    import docx  # python-docx
    doc = docx.Document(io.BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())
    return "\n".join(paragraphs)


def extract_pptx(data: bytes) -> str:
    """Extract text from a .pptx binary blob, slide by slide."""
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    lines: List[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        slide_texts: List[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        slide_texts.append(text)
        if slide_texts:
            lines.append(f"[شريحة {i}]")
            lines.extend(slide_texts)
    return "\n".join(lines)


def extract_pdf(data: bytes) -> str:
    """Extract text from a .pdf binary blob using PyMuPDF."""
    import fitz  # pymupdf
    doc = fitz.open(stream=data, filetype="pdf")
    pages: List[str] = []
    for page in doc:
        text = page.get_text().strip()
        if text:
            pages.append(text)
    return "\n".join(pages)


def extract_txt(data: bytes) -> str:
    """Decode a plain-text file (tries UTF-8 then cp1256 for Arabic)."""
    for enc in ("utf-8", "utf-8-sig", "cp1256", "latin-1"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace")


# ─────────────────────────────────────────────
# Public API
# ─────────────────────────────────────────────

def extract_text_from_file(filename: str, data: bytes) -> str:
    """
    Route to the right extractor based on file extension.
    Returns extracted text or raises ValueError if format unsupported.
    """
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "docx":
        return extract_docx(data)
    elif ext == "pptx":
        return extract_pptx(data)
    elif ext == "pdf":
        return extract_pdf(data)
    elif ext in ("txt", "text"):
        return extract_txt(data)
    else:
        raise ValueError(f"صيغة الملف غير مدعومة: .{ext}  (المدعوم: docx, pptx, pdf, txt)")


def merge_files(files: List[Tuple[str, bytes]]) -> str:
    """
    Extract and merge text from multiple files.
    Each file block is separated by a clear header:
      ════════════════════
      📄 اسم_الملف.docx
      ════════════════════
      <text>

    Args:
        files: list of (filename, bytes) tuples

    Returns:
        Merged text string (all files combined).
    """
    parts: List[str] = []
    for filename, data in files:
        try:
            text = extract_text_from_file(filename, data).strip()
            if not text:
                logger.warning("[FileExtractor] ملف فارغ بعد الاستخراج: %s", filename)
                text = "(لم يُستخرج أي نص من هذا الملف)"
        except Exception as exc:
            logger.error("[FileExtractor] خطأ في استخراج %s: %s", filename, exc)
            text = f"(تعذّر استخراج النص من هذا الملف: {exc})"

        separator = "═" * 50
        parts.append(
            f"{separator}\n📄 {filename}\n{separator}\n{text}"
        )

    return "\n\n".join(parts)
