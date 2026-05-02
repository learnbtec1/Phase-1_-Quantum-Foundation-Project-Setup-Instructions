# -*- coding: utf-8 -*-
"""
Multi-file student submission: extract PDF / DOCX / PPTX (and raw txt/md), merge into one document.
Each file is decoded and extracted independently; failures are logged and other files still merge.
"""
from __future__ import annotations

import logging
import re
from io import BytesIO
from pathlib import Path
from typing import List, Optional, Tuple

from docx import Document
from pypdf import PdfReader
from pptx import Presentation

logger = logging.getLogger(__name__)

FILE_BREAK = "\n\n---FILE BREAK---\n\n"

# Marks each part so the grader can see which file content came from (distribution / context).
TYPED_SUBMISSION_LABEL = "inline_submission"  # pseudo-name for pasted/typed work when no uploads

# Same intent as Next /api/parse-file
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md"}

MAX_BYTES_PER_FILE = 12 * 1024 * 1024
MAX_BYTES_TOTAL = 25 * 1024 * 1024

# Optional: more robust PDF text (installed via requirements)
try:
    import fitz  # PyMuPDF

    _HAS_PYMUPDF = True
except Exception:  # noqa: BLE001
    _HAS_PYMUPDF = False
    fitz = None  # type: ignore[misc, assignment]


def _extract_pdf_pymupdf(data: bytes) -> Optional[str]:
    if not _HAS_PYMUPDF or not data:
        return None
    try:
        doc = fitz.open(stream=data, filetype="pdf")
        try:
            parts: List[str] = []
            for i in range(len(doc)):
                try:
                    parts.append((doc[i].get_text() or "").strip())
                except Exception as e:  # noqa: BLE001
                    logger.warning("submission: pdf page extract failed (pymupdf) page=%s: %s", i, e)
            text = "\n".join(p for p in parts if p)
        finally:
            doc.close()
        return text or None
    except Exception as e:  # noqa: BLE001
        logger.warning("submission: pdf pymupdf failed, will try pypdf: %s", e)
        return None


def _extract_pdf_pypdf(data: bytes) -> str:
    if not data:
        return ""
    bio = BytesIO(data)
    try:
        reader = PdfReader(bio)
        out: List[str] = []
        for i, page in enumerate(reader.pages):
            try:
                out.append((page.extract_text() or "").strip())
            except Exception as e:  # noqa: BLE001
                logger.warning("submission: pdf page extract failed (pypdf) page=%s: %s", i, e)
        return "\n".join(p for p in out if p)
    except Exception as e:  # noqa: BLE001
        logger.error("submission: pdf pypdf failed: %s", e)
        return ""


def _extract_pdf_text(data: bytes) -> str:
    primary = _extract_pdf_pymupdf(data)
    if primary and primary.strip():
        return primary
    secondary = _extract_pdf_pypdf(data)
    if secondary and secondary.strip():
        return secondary
    if primary is not None:
        return (primary or "").strip()
    return (secondary or "").strip()


def _extract_docx_text(data: bytes) -> str:
    if not data:
        return ""
    bio = BytesIO(data)
    try:
        doc = Document(bio)
        out: List[str] = []
        for p in doc.paragraphs:
            if p.text and p.text.strip():
                out.append(p.text)
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    out.append(" | ".join(cells))
        return "\n".join(out)
    except Exception as e:  # noqa: BLE001
        logger.error("submission: docx extract failed: %s", e)
        return ""


def _extract_pptx_text(data: bytes) -> str:
    if not data:
        return ""
    bio = BytesIO(data)
    try:
        prs = Presentation(bio)
        parts: List[str] = []
        for sn, slide in enumerate(prs.slides):
            try:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        parts.append(shape.text)
                if slide.has_notes_slide:
                    parts.append(slide.notes_slide.notes_text_frame.text)
            except Exception as e:  # noqa: BLE001
                logger.warning("submission: pptx slide %s failed: %s", sn, e)
        return "\n".join(parts)
    except Exception as e:  # noqa: BLE001
        logger.error("submission: pptx extract failed: %s", e)
        return ""


def extract_text_from_bytes(filename: str, data: bytes) -> str:
    """
    Best-effort text for one file. Does not raise; returns "" on failure.
    """
    if not data:
        return ""
    ext = Path(filename or "").suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        logger.warning("submission: unsupported file type: %s", filename)
        return ""
    try:
        if ext == ".pdf":
            return _extract_pdf_text(data)
        if ext == ".docx":
            return _extract_docx_text(data)
        if ext == ".pptx":
            return _extract_pptx_text(data)
        if ext in (".txt", ".md"):
            return data.decode("utf-8", errors="replace")
    except Exception as e:  # noqa: BLE001
        logger.error("submission: extract failed for %s: %s", filename, e, exc_info=True)
    return ""


def _safe_display_name(filename: str) -> str:
    s = (filename or "unnamed").replace("\n", " ").replace("\r", " ")
    s = s.replace("]", "_")  # do not break [FILE: ...] when echoed in models
    return s.strip()[:300] or "unnamed"


def _file_aware_block(display_name: str, body: str) -> str:
    """[FILE: name]\\n...text... — one logical document part for the model."""
    b = (body or "").strip()
    if not b:
        return ""
    return f"[FILE: {_safe_display_name(display_name)}]\n{b}"


def build_combined_student_work(
    *,
    file_parts: List[Tuple[str, bytes]],
    student_work: str = "",
) -> str:
    """
    Merge all file texts in `file_parts` order (with [FILE: ...] markers), then append typed
    `student_work` (if any) with its own marker. Between parts: FILE_BREAK.
    One unified string for PASS1 / grader (PASS2) and RAG.
    """
    merged, _recv, _proc = _build_combined_with_metrics(
        file_parts=file_parts, student_work=student_work
    )
    return merged


def _build_combined_with_metrics(
    *,
    file_parts: List[Tuple[str, bytes]],
    student_work: str = "",
) -> tuple[str, int, int]:
    """
    Returns (merged_text, files_received, files_with_extracted_text).
    """
    file_texts: List[str] = []
    n_received = len(file_parts)
    n_processed = 0
    for name, raw in file_parts:
        try:
            t = extract_text_from_bytes(name, raw).strip()
        except Exception as e:  # noqa: BLE001
            logger.error("submission: unexpected error for %s: %s", name, e, exc_info=True)
            continue
        if t:
            n_processed += 1
            block = _file_aware_block(name, t)
            if block:
                file_texts.append(block)
        else:
            logger.warning("submission: no text extracted from file (skipping block): %s", name)
    merged_files = FILE_BREAK.join(file_texts) if file_texts else ""
    typed = (student_work or "").strip()
    typed_block = _file_aware_block(TYPED_SUBMISSION_LABEL, typed) if typed else ""

    if merged_files and typed_block:
        merged = merged_files + FILE_BREAK + typed_block
    elif merged_files:
        merged = merged_files
    elif typed_block:
        merged = typed_block
    else:
        merged = ""
    return merged, n_received, n_processed


def merge_and_log(
    *,
    file_parts: List[Tuple[str, bytes]],
    student_work: str = "",
) -> str:
    combined, n_recv, n_proc = _build_combined_with_metrics(
        file_parts=file_parts, student_work=student_work
    )
    logger.info(f"[FILES RECEIVED] {n_recv}")
    logger.info(f"[FILES PROCESSED] {n_proc}")
    logger.info("[FILES] total_files=%d", n_recv)
    logger.info("[FILES] total_chars=%d", len(combined))
    return combined


def parse_submission_file_sections(merged: str) -> List[Tuple[str, str]]:
    """
    Parse merged student work into (display_name, body) for each [FILE: name] part.
    Order is preserved. Body does not include the [FILE: ...] line.
    If there are no [FILE: ...] markers, one section ("submission", full_text) is returned.
    """
    s = (merged or "").strip()
    if not s:
        return []
    parts = s.split(FILE_BREAK)
    out: List[Tuple[str, str]] = []
    for p in parts:
        p = p.strip()
        if not p:
            continue
        m = re.match(r"^\[FILE: ([^\]]+)\]\s*\n?(.*)\Z", p, re.DOTALL)
        if not m:
            continue
        name, body = m.group(1).strip(), m.group(2) or ""
        if not name:
            name = "unnamed"
        out.append((name, body))
    if not out:
        return [("submission", s)]
    return out


def list_submission_file_names(merged: str) -> List[str]:
    """Unique logical source labels from [FILE: ...] markers, in first-seen order."""
    seen: set[str] = set()
    ordered: List[str] = []
    for name, _ in parse_submission_file_sections(merged):
        if name not in seen:
            seen.add(name)
            ordered.append(name)
    return ordered


def _norm_fn(name: str) -> str:
    return (name or "").replace("\\", "/").split("/")[-1].strip().lower()


def resolve_section_name(claimed: str, sections: List[Tuple[str, str]]) -> Optional[str]:
    """
    Map a model-provided file label to a canonical section name from the merged submission.
    """
    c = (claimed or "").strip()
    if not c:
        return None
    c_n = _norm_fn(c)
    for real, _ in sections:
        if _norm_fn(real) == c_n or real == c:
            return real
    # suffix / partial match (best effort)
    for real, _ in sections:
        if c_n and (_norm_fn(real).endswith(c_n) or c_n.endswith(_norm_fn(real))):
            return real
    return None


def infer_source_file_for_quote(quote: str, sections: List[Tuple[str, str]]) -> Optional[str]:
    """Which [FILE: ...] section first contains this verbatim quote (if any)."""
    q = (quote or "").strip()
    if len(q) < 2:
        return None
    lo = q.lower()
    for name, body in sections:
        b = body or ""
        if q in b:
            return name
        blo = b.lower()
        if lo in blo:
            try:
                idx = blo.find(lo)
            except Exception:
                continue
            if idx >= 0:
                return name
    return None
