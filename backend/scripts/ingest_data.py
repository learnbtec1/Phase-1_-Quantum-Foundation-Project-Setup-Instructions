# -*- coding: utf-8 -*-
"""
تطوير: كوجني - معلّم BTEC الرقمي (Eduverse-1)
الوصف: حقن شامل ومتشعب للمناهج (PDF, DOCX, PPTX, TXT)
المستهدف: L2, L3, Grade 10-12 و btec_specs
"""
from __future__ import annotations

import os
import sys
import logging
import argparse
from pathlib import Path
from typing import Iterable, List

# 1. إعداد المسارات لضمان رؤية موديولات المشروع
_BACK_END = Path(__file__).resolve().parent.parent
if str(_BACK_END) not in sys.path:
    sys.path.insert(0, str(_BACK_END))

# 2. تحميل البيئة يدوياً لتجاوز خطأ الـ Settings
from dotenv import load_dotenv
load_dotenv()

# 3. الاستيرادات التقنية
from docx import Document
from pptx import Presentation
from pypdf import PdfReader

from app.core.config import settings
from app.services.academic_rag_context import infer_grade_tier, infer_subject_key
from app.services.rag_documents_service import (
    TIER_TO_GRADE_NUM,
    extract_metadata_from_path,
    get_rag_documents_service,
)
from app.services.vector_service import get_vector_service

# إعداد الـ Logging
logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
logger = logging.getLogger(__name__)

# الفلاتر والامتدادات
EXTENSIONS = {
    ".pdf": "PDF",
    ".docx": "Word",
    ".pptx": "PowerPoint",
    ".txt": "Text",
    ".md": "Markdown"
}

def chunk_text_by_words(text: str, chunk_words: int = 500, overlap_words: int = 50) -> List[str]:
    """تقسيم النص لقطع معرفية مع الحفاظ على التداخل (Overlap)"""
    words = (text or "").split()
    if not words: return []
    step = max(1, chunk_words - overlap_words)
    chunks = []
    for i in range(0, len(words), step):
        piece = words[i : i + chunk_words]
        if piece:
            chunks.append(" ".join(piece))
    return chunks

def extract_text(file_path: Path) -> str:
    """استخراج النص بناءً على نوع الملف مع دعم الترميز العربي"""
    ext = file_path.suffix.lower()
    try:
        if ext == ".pdf":
            reader = PdfReader(str(file_path))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        
        if ext == ".docx":
            doc = Document(str(file_path))
            return "\n".join(p.text for p in doc.paragraphs if p.text)
        
        if ext == ".pptx":
            prs = Presentation(str(file_path))
            parts = []
            for slide in prs.slides:
                parts.extend([s.text for s in slide.shapes if hasattr(s, "text") and s.text])
                if slide.has_notes_slide:
                    parts.append(slide.notes_slide.notes_text_frame.text)
            return "\n".join(parts)
        
        if ext in [".txt", ".md"]:
            return file_path.read_text(encoding="utf-8", errors="replace")
            
    except Exception as e:
        logger.warning(f"⚠️ فشل قراءة الملف {file_path.name}: {e}")
    return ""

def iter_all_files(root_dir: Path) -> Iterable[Path]:
    """البحث المتشعب (Recursive) في كل المجلدات الفرعية مهما كان عمقها"""
    if not root_dir.exists():
        logger.error(f"❌ المجلد غير موجود: {root_dir}")
        return
    
    for root, _, files in os.walk(root_dir):
        for name in files:
            p = Path(root) / name
            if p.suffix.lower() in EXTENSIONS:
                yield p

def main():
    # قراءة المسار من .env أو استخدام الافتراضي
    raw_path = os.getenv("LOCAL_RAG_DIR") or str(_BACK_END / "data")
    root_path = Path(raw_path).resolve()
    
    print(f"\n🚀 جاري بدء عملية الحقن المعرفي...")
    print(f"📂 المجلد الرئيسي: {root_path}")
    
    if not settings.OPENAI_API_KEY:
        logger.error("🚨 خطأ: OPENAI_API_KEY غير موجود في ملف .env")
        return

    svc = get_vector_service()
    svc.ensure_schema()
    rsvc = get_rag_documents_service()
    rsvc.ensure_schema()

    total_files = 0
    total_chunks = 0
    total_rag_rows = 0

    # البدء بالمسح المتشعب
    all_files = list(iter_all_files(root_path))
    print(f"📊 تم العثور على {len(all_files)} ملف متوافق.\n")

    for path in all_files:
        text = extract_text(path)
        if not text.strip(): continue
        
        chunks = chunk_text_by_words(text)
        if not chunks: continue
        
        # إعداد الميتا-داتا (مهم للبحث لاحقاً)
        source_id = str(path.resolve())
        meta = {
            "source_file": path.name,
            "source_path": source_id,
            "category": path.parent.name,  # اسم المجلد الفرعي (مثل Unit 4)
        }

        path_key = str(path)
        pe = extract_metadata_from_path(path)
        tier = infer_grade_tier(path_key)
        g = TIER_TO_GRADE_NUM.get(tier) or pe.get("grade")
        subj = pe.get("subject")
        if not subj:
            sk = infer_subject_key(path.parent.name) or infer_subject_key(path.stem)
            if sk and sk != "general":
                subj = sk
        path_extras = {**pe, "grade": g, "subject": subj or pe.get("subject")}

        # تنظيف البيانات القديمة لنفس الملف لتجنب التكرار
        svc.delete_by_source_id(source_id)

        # الحقن في pgvector (JSONB)
        ids = svc.ingest_chunks(chunks, metadata=meta, source_id=source_id)

        # نفس الملف في rag_documents (أعمدة + فلترة SQL)
        n_rag = rsvc.ingest_file_chunks(
            path,
            chunks,
            path_extras=path_extras,
        )
        total_rag_rows += n_rag

        total_files += 1
        total_chunks += len(ids)
        logger.info(
            f"✅ تم حقن: {path.name} ({len(ids)} Chunks, rag_documents={n_rag})"
        )

    print(f"\n✨ تمت المهمة بنجاح!")
    print(f"📚 إجمالي الملفات: {total_files}")
    print(f"🧩 إجمالي الـ Chunks: {total_chunks}")
    print(f"📇 إجمالي صفوف rag_documents: {total_rag_rows}")

if __name__ == "__main__":
    main()